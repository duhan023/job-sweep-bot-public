#!/usr/bin/env python3
"""
configure.py - Interactive onboarding for Job Sweep Bot.

Run once after forking:

    python configure.py

Produces:
    config.json                          (all user settings; no code edits ever needed)
    areas/job-search-log.md              (dedup ledger)
    .github/workflows/daily_sweep.yml    (cron schedule in UTC, derived from your local time)

Works with no third-party packages installed. pypdf / python-docx are used only
if present; otherwise you can paste resume text directly.
"""

from __future__ import annotations

import json
import os
import re
import sys
from collections import Counter
from datetime import date, datetime, time as dtime, timedelta, timezone
from pathlib import Path

try:
    from zoneinfo import ZoneInfo  # Python 3.9+
except ImportError:  # pragma: no cover
    ZoneInfo = None  # type: ignore

ROOT = Path(__file__).resolve().parent
CONFIG_PATH = ROOT / "config.json"
LOG_PATH = ROOT / "areas" / "job-search-log.md"
WORKFLOW_PATH = ROOT / ".github" / "workflows" / "daily_sweep.yml"

CONFIG_VERSION = 1

# --------------------------------------------------------------------------
# Terminal helpers
# --------------------------------------------------------------------------

BOLD = "\033[1m"
DIM = "\033[2m"
CYAN = "\033[36m"
GREEN = "\033[32m"
YELLOW = "\033[33m"
RED = "\033[31m"
RESET = "\033[0m"


def supports_color() -> bool:
    return sys.stdout.isatty() and os.environ.get("NO_COLOR") is None


def c(text: str, color: str) -> str:
    return f"{color}{text}{RESET}" if supports_color() else text


def header(text: str) -> None:
    line = "=" * 68
    print("\n" + c(line, DIM))
    print(c(text, BOLD))
    print(c(line, DIM))


def ask(prompt: str, default: str | None = None, required: bool = True) -> str:
    suffix = f" [{default}]" if default else ""
    while True:
        raw = input(c(f"{prompt}{suffix}: ", CYAN)).strip()
        if not raw and default is not None:
            return default
        if raw:
            return raw
        if not required:
            return ""
        print(c("  A value is required.", YELLOW))


def ask_int(prompt: str, default: int, lo: int, hi: int) -> int:
    while True:
        raw = ask(prompt, str(default))
        try:
            val = int(re.sub(r"[^0-9-]", "", raw))
        except ValueError:
            print(c("  Enter a whole number.", YELLOW))
            continue
        if lo <= val <= hi:
            return val
        print(c(f"  Enter a number between {lo} and {hi}.", YELLOW))


def ask_yes_no(prompt: str, default: bool = True) -> bool:
    d = "Y/n" if default else "y/N"
    while True:
        raw = input(c(f"{prompt} [{d}]: ", CYAN)).strip().lower()
        if not raw:
            return default
        if raw in ("y", "yes"):
            return True
        if raw in ("n", "no"):
            return False
        print(c("  Answer y or n.", YELLOW))


def ask_list(prompt: str, default: str | None = None) -> list[str]:
    raw = ask(prompt, default)
    return [p.strip() for p in raw.split(",") if p.strip()]


# --------------------------------------------------------------------------
# Resume ingestion
# --------------------------------------------------------------------------

def read_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        try:
            from PyPDF2 import PdfReader  # type: ignore
        except ImportError:
            raise RuntimeError(
                "pypdf is not installed. Run `pip install pypdf`, or export your "
                "resume to .txt and re-run."
            )
    reader = PdfReader(str(path))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def read_docx(path: Path) -> str:
    try:
        import docx  # python-docx
    except ImportError:
        raise RuntimeError(
            "python-docx is not installed. Run `pip install python-docx`, or save "
            "your resume as .txt and re-run."
        )
    document = docx.Document(str(path))
    chunks = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            chunks.extend(cell.text for cell in row.cells)
    return "\n".join(chunks)


def read_txt(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


def load_resume_text() -> str:
    print(
        "\nPoint me at your master resume. Supported: .pdf, .docx, .txt\n"
        "Or press Enter with no path to paste the text instead."
    )
    while True:
        raw = input(c("Resume file path (blank to paste): ", CYAN)).strip().strip('"').strip("'")
        if not raw:
            return paste_resume_text()
        path = Path(os.path.expanduser(raw)).resolve()
        if not path.exists():
            print(c(f"  Not found: {path}", YELLOW))
            continue
        ext = path.suffix.lower()
        try:
            if ext == ".pdf":
                text = read_pdf(path)
            elif ext in (".docx", ".doc"):
                text = read_docx(path)
            else:
                text = read_txt(path)
        except Exception as exc:  # noqa: BLE001
            print(c(f"  Could not read that file: {exc}", RED))
            continue
        if len(text.strip()) < 200:
            print(
                c(
                    "  Extracted almost no text (likely a scanned/image PDF). "
                    "Paste the text instead.",
                    YELLOW,
                )
            )
            return paste_resume_text()
        print(c(f"  Read {len(text.split())} words from {path.name}.", GREEN))
        return text


def paste_resume_text() -> str:
    print(
        "\nPaste your resume text below. When finished, type END on its own line "
        "and press Enter."
    )
    lines: list[str] = []
    while True:
        try:
            line = input()
        except EOFError:
            break
        if line.strip().upper() == "END":
            break
        lines.append(line)
    text = "\n".join(lines)
    while len(text.strip()) < 100:
        print(c("That is too short to score against. Paste more of the resume.", YELLOW))
        return paste_resume_text()
    return text


# --------------------------------------------------------------------------
# Keyword extraction (industry-agnostic, frequency + section weighted)
# --------------------------------------------------------------------------

STOPWORDS = set(
    """
a about above across after again against all also am an and any are aren as at be because been
before being below between both but by can cannot could couldn did didn do does doesn doing don
down during each few for from further had hadn has hasn have haven having he her here hers herself
him himself his how i if in into is isn it its itself just ll me more most must my myself no nor
not now of off on once only or other ought our ours ourselves out over own re s same shan she
should shouldn so some such t than that the their theirs them themselves then there these they
this those through to too under until up ve very was wasn we were weren what when where which
while who whom why will with won would wouldn you your yours yourself yourselves
""".split()
)

# Resume boilerplate that carries no matching signal.
RESUME_NOISE = set(
    """
ability able accomplished achieve achieved achievements across activities additional address
adept advanced ahead align aligned aligning among amount analysis analyze and/or annual applied
approach approaches areas around assist assisted assisting associate attention award awarded based
basic best better broad build building built career certification certifications certified charge
clearly close collaborate collaborated collaboration colleagues company complete completed complex
comprehensive conduct conducted consistently contributed coordinate coordinated core created
creating cross cross-functional current currently daily date day deadlines degree deliver
delivered delivering delivery demonstrated description designed detail details develop developed
developing different direct directed director drive driven drove duties education effective
efficiency efficient effort employee employees employment enable enabled end ensure ensured
ensuring environment established etc excellent execute executed executing execution existing
expected experience experienced expertise external facilitate facilitated familiar first focus
focused following full function functional gained general goal goals graduate great group grow
growth hands help helped high highly hire history hours implement implemented implementing improve
improved improvement improvements include included includes including increase increased
individual industry initiative initiatives inside internal issues job join key knowledge large
lead leading learn learned led level leverage leveraged liaison like located location maintain
maintained maintaining major manage managed management manager managing many meet meeting meetings
member members methods monthly month multiple name nbsp need needs new number objectives obtained
office ongoing open operations opportunities order organization organizational other overall
oversaw oversee overseeing part participate participated partner partners performance performed
period person personal phone plan planned planning plans point position possible practices prepare
prepared present presented previous prior proactively problem problems procedures process
processes produced professional proficient program programs progress project projects proven
provide provided providing quality quarterly range rate reduce reduced reducing references
regarding related relevant reported reporting requirements requires research resolve resolved
responsibilities responsible result results resume review reviewed reviews role roles run running
ran raise raised raising quarter per own owned serving cutting migrating migrated driving drove drives
schedule scheduled scope seeking senior serve served service services set several skill skills
solution solutions specific staff standard standards start started state status strong subject
success successful successfully summary supervised support supported supporting system systems
take target task tasks team teams technical technologies technology throughout time timely tools
top total track tracking train trained training two university update updated updates use used
using utilized utilizing value various via week weekly well within work worked working works
year years
""".split()
)

DROP = STOPWORDS | RESUME_NOISE

SECTION_HEADER_RE = re.compile(
    r"^\s*(technical\s+skills|core\s+competenc\w*|skills?|tools?\s*&?\s*technolog\w*|"
    r"technolog\w*|certification[s]?|licen[cs]es?|proficienc\w*|expertise|"
    r"key\s+skills|areas\s+of\s+expertise|software|platforms?)\s*[:\-]?\s*$",
    re.IGNORECASE,
)

TOKEN_RE = re.compile(r"[A-Za-z][A-Za-z0-9+#./&-]{1,29}")


def _clean_token(tok: str) -> str:
    tok = tok.strip(".,;:/&-").lower()
    return tok


def extract_skill_sections(text: str) -> str:
    """Return the text of resume sections that look like skill inventories."""
    lines = text.splitlines()
    captured: list[str] = []
    capturing = False
    blanks = 0
    for line in lines:
        stripped = line.strip()
        if SECTION_HEADER_RE.match(stripped):
            capturing = True
            blanks = 0
            continue
        if capturing:
            if not stripped:
                blanks += 1
                if blanks >= 2:
                    capturing = False
                continue
            # A new all-caps heading ends the section.
            if len(stripped) < 40 and stripped.isupper() and not any(ch.isdigit() for ch in stripped):
                capturing = False
                continue
            captured.append(stripped)
            blanks = 0
    return "\n".join(captured)


SEGMENT_SPLIT_RE = re.compile(r"[,;:|•·\n\r\t()\[\]{}]+|(?<!\d)\.(?!\d)| - | / ")


def tokenize(text: str) -> list[str]:
    return [_clean_token(t) for t in TOKEN_RE.findall(text)]


def segment_tokens(text: str) -> list[list[str]]:
    """
    Tokenize, but keep punctuation boundaries so bigrams never straddle them.
    'Excel, SQL, Power BI' yields [[excel],[sql],[power, bi]] -> only 'power bi'
    becomes a candidate phrase.
    """
    return [tokenize(seg) for seg in SEGMENT_SPLIT_RE.split(text) if seg and seg.strip()]


EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.]+")
URL_RE = re.compile(
    r"https?://\S+|\bwww\.\S+|\b[\w-]+\.(?:com|io|net|org|dev|co|ai|me|edu)\b\S*",
    re.IGNORECASE,
)
CITY_STATE_RE = re.compile(r"^[A-Za-z .'-]{2,30},\s*[A-Za-z]{2,}\b")
PHONE_RE = re.compile(r"\(?\d{3}\)?[\s.-]?\d{3}[\s.-]?\d{4}")


def strip_contact_header(text: str) -> str:
    """
    Remove the name/contact block. Those lines produce phrases like
    'marcus reed' that match nothing in a job description.
    """
    text = EMAIL_RE.sub(" ", text)
    text = URL_RE.sub(" ", text)
    text = PHONE_RE.sub(" ", text)
    lines = text.splitlines()
    blanked_name = False
    for i, line in enumerate(lines[:4]):
        stripped = line.strip()
        if not stripped:
            continue
        if not blanked_name:
            lines[i] = ""  # first non-empty line is almost always the name
            blanked_name = True
            continue
        # Contact/location lines: short, pipe-separated, or "City, ST"
        if len(stripped) < 70 and ("|" in stripped or CITY_STATE_RE.match(stripped)):
            lines[i] = ""
    return "\n".join(lines)


def extract_keywords(resume_text: str, top_n: int = 60) -> list[str]:
    """
    Frequency-and-position weighted keyword extraction.

    Unigrams and phrases are scored. Terms inside an explicit skills / tools /
    certifications section get a 3x boost, because that is where users put the
    vocabulary a job description will actually echo back. Phrases pulled from
    prose must appear at least twice to survive, which removes one-off sentence
    fragments like "owned demand" while keeping real terms like "demand
    generation".
    """
    cleaned = strip_contact_header(resume_text)
    skill_text = extract_skill_sections(resume_text)
    body_segments = segment_tokens(cleaned)
    skill_segments = segment_tokens(skill_text)

    scores: Counter[str] = Counter()
    body_phrase_counts: Counter[str] = Counter()
    skill_terms: set[str] = set()

    def usable(tok: str) -> bool:
        return tok not in DROP and len(tok) >= 2 and not tok.isdigit() and not tok.isnumeric()

    def add_grams(segments: list[list[str]], weight: float, is_skill_section: bool) -> None:
        for tokens in segments:
            for tok in tokens:
                if not usable(tok):
                    continue
                scores[tok] += weight
                if is_skill_section:
                    skill_terms.add(tok)
            for a, b in zip(tokens, tokens[1:]):
                if not usable(a) or not usable(b):
                    continue
                phrase = f"{a} {b}"
                scores[phrase] += weight * 1.4  # phrases are higher-signal
                if is_skill_section:
                    skill_terms.add(phrase)
                else:
                    body_phrase_counts[phrase] += 1

    add_grams(body_segments, 1.0, False)
    add_grams(skill_segments, 3.0, True)

    # Prose phrases that appear once are sentence fragments, not vocabulary.
    for phrase, count in body_phrase_counts.items():
        if count < 2 and phrase not in skill_terms:
            scores.pop(phrase, None)

    # Drop unigrams that only ever appear inside a stronger phrase.
    ranked = [term for term, _ in scores.most_common(top_n * 3)]
    phrases = [t for t in ranked if " " in t]
    result: list[str] = []
    for term in ranked:
        if " " not in term:
            subsumed = any(term in p.split() for p in phrases[:top_n])
            if subsumed and scores[term] < 4:
                continue
        result.append(term)
        if len(result) >= top_n:
            break
    return result


def review_keywords(keywords: list[str]) -> list[str]:
    while True:
        print("\n" + c("Extracted resume keywords (used for Fit Scoring):", BOLD))
        for i in range(0, len(keywords), 3):
            row = keywords[i : i + 3]
            print("  " + "".join(f"{i + j + 1:>3}. {t:<26}" for j, t in enumerate(row)))
        print(
            c(
                "\nEnter numbers to REMOVE (e.g. 3,7,12), '+term1, term2' to ADD, "
                "or press Enter to accept.",
                DIM,
            )
        )
        raw = input(c("> ", CYAN)).strip()
        if not raw:
            return keywords
        if raw.startswith("+"):
            for term in raw[1:].split(","):
                term = term.strip().lower()
                if term and term not in keywords:
                    keywords.append(term)
            continue
        try:
            drop_idx = {int(x.strip()) - 1 for x in raw.split(",") if x.strip()}
        except ValueError:
            print(c("  Could not parse that. Use numbers, or +term.", YELLOW))
            continue
        keywords = [t for i, t in enumerate(keywords) if i not in drop_idx]


# --------------------------------------------------------------------------
# Schedule -> cron
# --------------------------------------------------------------------------

TZ_ALIASES = {
    "eastern": "America/New_York", "et": "America/New_York", "est": "America/New_York",
    "edt": "America/New_York", "new york": "America/New_York", "nyc": "America/New_York",
    "central": "America/Chicago", "ct": "America/Chicago", "cst": "America/Chicago",
    "cdt": "America/Chicago", "chicago": "America/Chicago", "dallas": "America/Chicago",
    "mountain": "America/Denver", "mt": "America/Denver", "mst": "America/Denver",
    "mdt": "America/Denver", "denver": "America/Denver",
    "arizona": "America/Phoenix", "phoenix": "America/Phoenix",
    "pacific": "America/Los_Angeles", "pt": "America/Los_Angeles",
    "pst": "America/Los_Angeles", "pdt": "America/Los_Angeles",
    "los angeles": "America/Los_Angeles", "seattle": "America/Los_Angeles",
    "alaska": "America/Anchorage", "hawaii": "Pacific/Honolulu",
    "utc": "UTC", "gmt": "UTC", "zulu": "UTC",
    "uk": "Europe/London", "london": "Europe/London", "bst": "Europe/London",
    "ireland": "Europe/Dublin", "dublin": "Europe/Dublin",
    "cet": "Europe/Berlin", "cest": "Europe/Berlin", "berlin": "Europe/Berlin",
    "paris": "Europe/Paris", "amsterdam": "Europe/Amsterdam", "madrid": "Europe/Madrid",
    "warsaw": "Europe/Warsaw", "stockholm": "Europe/Stockholm",
    "ist": "Asia/Kolkata", "india": "Asia/Kolkata", "kolkata": "Asia/Kolkata",
    "bangalore": "Asia/Kolkata", "mumbai": "Asia/Kolkata", "delhi": "Asia/Kolkata",
    "dubai": "Asia/Dubai", "gulf": "Asia/Dubai",
    "singapore": "Asia/Singapore", "sgt": "Asia/Singapore",
    "manila": "Asia/Manila", "philippines": "Asia/Manila",
    "tokyo": "Asia/Tokyo", "jst": "Asia/Tokyo", "japan": "Asia/Tokyo",
    "sydney": "Australia/Sydney", "aest": "Australia/Sydney",
    "melbourne": "Australia/Melbourne", "brisbane": "Australia/Brisbane",
    "perth": "Australia/Perth", "auckland": "Pacific/Auckland",
    "toronto": "America/Toronto", "vancouver": "America/Vancouver",
    "sao paulo": "America/Sao_Paulo", "brazil": "America/Sao_Paulo",
    "mexico": "America/Mexico_City", "mexico city": "America/Mexico_City",
    "johannesburg": "Africa/Johannesburg", "lagos": "Africa/Lagos",
    "nairobi": "Africa/Nairobi", "cairo": "Africa/Cairo",
}

TIME_RE = re.compile(
    r"^\s*(\d{1,2})(?::|\.)?(\d{2})?\s*(am|pm)?\s*(.*)$", re.IGNORECASE
)
COMPACT_TIME_RE = re.compile(r"^\s*(\d{3,4})\s*(am|pm)?\s*(.*)$", re.IGNORECASE)


def parse_time_and_zone(raw: str) -> tuple[int, int, str]:
    """
    Parse '08:30 AM Eastern', '8am IST', '17:45 Europe/Berlin', '0630 pacific'.
    Returns (hour_24, minute, iana_timezone).
    """
    compact = COMPACT_TIME_RE.match(raw)
    if compact:
        digits = compact.group(1).zfill(4)
        hour, minute = int(digits[:2]), int(digits[2:])
        meridiem = (compact.group(2) or "").lower()
        zone_raw = (compact.group(3) or "").strip()
    else:
        m = TIME_RE.match(raw)
        if not m:
            raise ValueError("Could not read a time out of that.")
        hour = int(m.group(1))
        minute = int(m.group(2) or 0)
        meridiem = (m.group(3) or "").lower()
        zone_raw = (m.group(4) or "").strip()

    if meridiem == "pm" and hour < 12:
        hour += 12
    if meridiem == "am" and hour == 12:
        hour = 0
    if not 0 <= hour <= 23 or not 0 <= minute <= 59:
        raise ValueError("Time out of range.")

    zone_key = zone_raw.strip(" ,.").lower()
    tz_name = None
    if not zone_key:
        tz_name = "UTC"
    elif "/" in zone_raw:
        tz_name = zone_raw
    elif zone_key in TZ_ALIASES:
        tz_name = TZ_ALIASES[zone_key]
    else:
        # strip trailing words like "time"
        zone_key2 = re.sub(r"\b(time|standard|daylight|zone)\b", "", zone_key).strip()
        tz_name = TZ_ALIASES.get(zone_key2)
    if not tz_name:
        raise ValueError(f"Unknown timezone: '{zone_raw}'. Try 'Eastern', 'IST', or 'Europe/Berlin'.")
    if ZoneInfo is not None:
        ZoneInfo(tz_name)  # raises if invalid
    return hour, minute, tz_name


def compress_dow(days: list[int]) -> str:
    """[1,2,3,4,5] -> '1-5' ; [2,3,4,5,6] -> '2-6' ; [0,6] -> '0,6'."""
    parts: list[str] = []
    start = prev = days[0]
    for day in days[1:] + [None]:  # type: ignore[list-item]
        if day is not None and day == prev + 1:
            prev = day
            continue
        if start == prev:
            parts.append(str(start))
        elif prev == start + 1:
            parts.extend([str(start), str(prev)])
        else:
            parts.append(f"{start}-{prev}")
        if day is not None:
            start = prev = day
    return ",".join(parts)


def local_schedule_to_cron(hour: int, minute: int, tz_name: str, weekdays_only: bool) -> tuple[str, str]:
    """
    Convert a local run time into a UTC cron expression.

    GitHub Actions cron is always UTC and has no DST awareness, so the offset is
    computed against the next occurrence and the caller is warned that the local
    run time will shift by one hour across a DST boundary.
    Returns (cron_expression, human_readable_utc_time).
    """
    if ZoneInfo is None:
        raise RuntimeError("Python 3.9+ with zoneinfo is required.")
    tz = ZoneInfo(tz_name)
    today = date.today()
    monday = today + timedelta(days=(7 - today.weekday()) % 7 or 7)  # next Monday

    local_days = range(5) if weekdays_only else range(7)
    utc_dows: set[int] = set()
    utc_hour = utc_minute = 0
    for offset in local_days:
        local_dt = datetime.combine(monday + timedelta(days=offset), dtime(hour, minute), tzinfo=tz)
        utc_dt = local_dt.astimezone(timezone.utc)
        utc_hour, utc_minute = utc_dt.hour, utc_dt.minute
        utc_dows.add((utc_dt.weekday() + 1) % 7)  # python Mon=0 -> cron Sun=0

    if len(utc_dows) == 7:
        dow = "*"
    else:
        dow = compress_dow(sorted(utc_dows))
    cron = f"{utc_minute} {utc_hour} * * {dow}"
    return cron, f"{utc_hour:02d}:{utc_minute:02d} UTC"


# --------------------------------------------------------------------------
# Salary parsing
# --------------------------------------------------------------------------

def parse_salary_floor(raw: str) -> int:
    """'$90,000' -> 90000 ; '90k' -> 90000 ; 'unlisted'/'' -> 0 (no floor)."""
    if not raw or raw.strip().lower() in ("unlisted", "none", "n/a", "no", "any", "0"):
        return 0
    cleaned = raw.lower().replace(",", "").replace("$", "").strip()
    hourly = bool(re.search(r"/\s*h|\bhour|\bhr\b", cleaned))
    m = re.search(r"(\d+(?:\.\d+)?)\s*([km])?", cleaned)
    if not m:
        return 0
    value = float(m.group(1))
    suffix = m.group(2)
    if suffix == "k":
        value *= 1_000
    elif suffix == "m":
        value *= 1_000_000
    elif hourly:
        value *= 2_080  # 40 hrs x 52 weeks
    elif value < 1_000:  # someone typed "90" meaning 90k
        value *= 1_000
    return int(value)


# Words people type to mean "I have no deal-breakers". Treated as an empty list,
# never as a literal substring to match against job descriptions.
NOOP_ANSWERS = {
    "nothing", "none", "no", "na", "n/a", "nil", "nope", "not applicable",
    "no preference", "any", "anything", "all", "-", "--", "x", "skip", "n a",
}


def normalize_exclusions(items: list[str]) -> list[str]:
    """'Exclude C2C' -> 'c2c' ; 'No Security Clearance' -> 'security clearance'.

    A no-op answer such as "nothing" or "none" yields an empty list. Treating it
    literally would silently drop every posting whose text happens to contain
    that word.
    """
    out: list[str] = []
    for item in items:
        term = item.strip().lower()
        if term in NOOP_ANSWERS:
            continue
        term = re.sub(r"^(exclude|no|not|avoid|skip|without)\s+", "", term)
        term = term.strip(" .,-")
        if not term or term in NOOP_ANSWERS:
            continue
        if len(term) < 2:
            continue
        out.append(term)
    return out


# --------------------------------------------------------------------------
# Asset writers
# --------------------------------------------------------------------------

WORKFLOW_TEMPLATE = """# Auto-generated by configure.py -- safe to regenerate at any time.
# Schedule below is UTC. Local equivalent: {local_desc}
name: Daily Job Sweep

on:
  schedule:
    - cron: "{cron}"
  workflow_dispatch:

permissions:
  contents: write

concurrency:
  group: job-sweep
  cancel-in-progress: false

jobs:
  sweep:
    runs-on: ubuntu-latest
    timeout-minutes: 20

    steps:
      - name: Checkout repository
        uses: actions/checkout@v4

      - name: Set up Python 3.11
        uses: actions/setup-python@v5
        with:
          python-version: "3.11"
          cache: pip

      - name: Install dependencies
        run: |
          python -m pip install --upgrade pip
          pip install -r requirements.txt

      - name: Run job sweep
        env:
          GMAIL_USER: ${{{{ secrets.GMAIL_USER }}}}
          GMAIL_APP_PASSWORD: ${{{{ secrets.GMAIL_APP_PASSWORD }}}}
        run: python sweep_public.py

      - name: Commit updated job log
        run: |
          git config user.name "github-actions[bot]"
          git config user.email "41898282+github-actions[bot]@users.noreply.github.com"
          git add areas/job-search-log.md
          if git diff --staged --quiet; then
            echo "No new roles to log."
          else
            git commit -m "chore: log dispatched roles ($(date -u +'%Y-%m-%d %H:%M UTC'))"
            git push
          fi
"""

LOG_TEMPLATE = """# Job Search Log

Append-only ledger written by `sweep_public.py`. Every role that has been emailed
to you is recorded here so it is never emailed twice. Do not delete rows unless
you want those roles to resurface.

| Date Sent | Company | Role | Location | Fit | Link | Key |
| --------- | ------- | ---- | -------- | --- | ---- | --- |
"""


def write_config(cfg: dict) -> None:
    CONFIG_PATH.write_text(json.dumps(cfg, indent=2) + "\n", encoding="utf-8")
    print(c(f"  wrote {CONFIG_PATH.relative_to(ROOT)}", GREEN))


def write_log() -> None:
    LOG_PATH.parent.mkdir(parents=True, exist_ok=True)
    if LOG_PATH.exists() and LOG_PATH.stat().st_size > 0:
        print(c(f"  kept existing {LOG_PATH.relative_to(ROOT)} (dedup history preserved)", DIM))
        return
    LOG_PATH.write_text(LOG_TEMPLATE, encoding="utf-8")
    print(c(f"  wrote {LOG_PATH.relative_to(ROOT)}", GREEN))


def write_workflow(cron: str, local_desc: str) -> None:
    WORKFLOW_PATH.parent.mkdir(parents=True, exist_ok=True)
    WORKFLOW_PATH.write_text(
        WORKFLOW_TEMPLATE.format(cron=cron, local_desc=local_desc), encoding="utf-8"
    )
    print(c(f"  wrote {WORKFLOW_PATH.relative_to(ROOT)}", GREEN))


# --------------------------------------------------------------------------
# Main flow
# --------------------------------------------------------------------------

def main() -> int:
    header("JOB SWEEP BOT - SETUP")
    print(
        "This takes about three minutes. Nothing you enter leaves your machine;\n"
        "it is written to config.json in this folder.\n"
    )

    existing = {}
    if CONFIG_PATH.exists():
        try:
            existing = json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
            print(c("Found an existing config.json - its values are the defaults below.\n", DIM))
        except json.JSONDecodeError:
            pass

    def prior(path: str, default):
        node = existing
        for part in path.split("."):
            if not isinstance(node, dict) or part not in node:
                return default
            node = node[part]
        return node

    # ---- 1. Resume ------------------------------------------------------
    header("STEP 1 of 4  -  Resume")
    resume_text = load_resume_text()
    keywords = extract_keywords(resume_text)
    keywords = review_keywords(keywords)

    # ---- 2. Targeting ---------------------------------------------------
    header("STEP 2 of 4  -  Targeting")
    print(
        "Titles drive the search queries. Use the exact phrasing employers post,\n"
        "not your internal title. 2-5 titles works best.\n"
    )
    titles = ask_list(
        'Target job titles (comma-separated)',
        ", ".join(prior("targeting.titles", [])) or None,
    )
    while not titles:
        titles = ask_list("At least one title is required")

    locations = ask_list(
        'Preferred locations (e.g. "Remote, Austin TX, London UK")',
        ", ".join(prior("targeting.locations", [])) or "Remote",
    )
    remote_ok = ask_yes_no("Include remote-flagged roles?", True)

    max_yoe = ask_int(
        "Max years of experience a posting may demand (excludes roles above your profile)",
        int(prior("targeting.max_years_experience", 5)),
        0,
        40,
    )
    salary_raw = ask(
        'Salary floor (e.g. "$90,000", "90k", or "Unlisted")',
        prior("targeting.salary_floor_raw", "Unlisted"),
    )
    salary_floor = parse_salary_floor(salary_raw)

    exclusions_raw = ask_list(
        'Deal-breakers (e.g. "Exclude C2C, No Security Clearance, Exclude Defense")',
        ", ".join(prior("targeting.exclusions_raw", [])) or "",
    )
    exclusions = normalize_exclusions(exclusions_raw)

    min_fit = ask_int(
        "Minimum Fit Score to email (1-100; 55 is a sane start)",
        int(prior("scoring.min_fit_score", 55)),
        1,
        100,
    )
    max_results = ask_int(
        "Max roles per email",
        int(prior("scoring.max_results", 25)),
        1,
        100,
    )

    print("\n" + c("Company ATS boards - this is your main source.", BOLD))
    print(
        c(
            "Greenhouse, Lever and Ashby are the employers' own job boards. Roles appear\n"
            "there before they reach any aggregator, and the APIs are public and stable.\n"
            "Format: type:token, comma separated. For example:\n"
            "  greenhouse:stripe, lever:netflix, ashby:ramp\n"
            "The token is the company slug in their careers URL:\n"
            "  boards.greenhouse.io/stripe  ->  greenhouse:stripe\n"
            "  jobs.lever.co/netflix        ->  lever:netflix\n"
            "Add 10-20 companies you actually want to work at. You can add more later\n"
            "by editing config.json.",
            DIM,
        )
    )
    prior_boards = prior("sources.ats_boards", [])
    prior_parts = []
    for b in prior_boards if isinstance(prior_boards, list) else []:
        if isinstance(b, dict) and b.get("type") and b.get("token"):
            prior_parts.append(f'{b["type"]}:{b["token"]}')
        elif isinstance(b, str) and ":" in b:
            prior_parts.append(b.strip())
    prior_boards_str = ", ".join(prior_parts)
    boards_raw = ask("ATS boards", prior_boards_str, required=False)
    ats_boards = []
    for chunk in boards_raw.split(","):
        chunk = chunk.strip()
        if not chunk or ":" not in chunk:
            continue
        kind, token = chunk.split(":", 1)
        kind = kind.strip().lower()
        if kind in ("greenhouse", "lever", "ashby"):
            ats_boards.append({"type": kind, "token": token.strip()})
        else:
            print(c(f"  Skipped '{chunk}' - type must be greenhouse, lever or ashby.", YELLOW))
    if ats_boards:
        print(c(f"  {len(ats_boards)} board(s) registered.", GREEN))
    else:
        print(
            c(
                "  No boards added. Google Careers alone only returns Google's own roles,\n"
                "  so your digest will be thin until you add some to config.json.",
                YELLOW,
            )
        )

    # ---- 3. Delivery ----------------------------------------------------
    header("STEP 3 of 4  -  Delivery")
    to_email = ask("Destination email address", prior("delivery.to_email", None))
    while "@" not in to_email:
        to_email = ask("That does not look like an email. Destination email address")

    while True:
        time_raw = ask(
            'Daily run time and timezone (e.g. "08:30 AM Eastern", "7am IST", "17:00 Europe/Berlin")',
            f'{prior("schedule.local_time", "08:30")} {prior("schedule.timezone", "America/New_York")}',
        )
        try:
            hour, minute, tz_name = parse_time_and_zone(time_raw)
            break
        except Exception as exc:  # noqa: BLE001
            print(c(f"  {exc}", YELLOW))

    weekdays_only = ask_yes_no("Weekdays only (Mon-Fri)? Answer n for every day", True)
    cron, utc_desc = local_schedule_to_cron(hour, minute, tz_name, weekdays_only)
    freq_label = "Mon-Fri" if weekdays_only else "every day"
    local_desc = f"{hour:02d}:{minute:02d} {tz_name}, {freq_label}"
    print(c(f"  Cron: {cron}  ({utc_desc}, {freq_label})", GREEN))
    print(
        c(
            "  Note: GitHub cron is UTC-only. Across a DST change your local run time\n"
            "  shifts by one hour until you re-run configure.py. Also, GitHub queues\n"
            "  scheduled jobs under load - expect 0-20 minutes of drift.",
            DIM,
        )
    )

    # ---- 4. Write assets ------------------------------------------------
    header("STEP 4 of 4  -  Writing files")
    cfg = {
        "version": CONFIG_VERSION,
        "generated_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "resume_keywords": keywords,
        "targeting": {
            "titles": titles,
            "locations": locations,
            "remote_ok": remote_ok,
            "max_years_experience": max_yoe,
            "salary_floor": salary_floor,
            "salary_floor_raw": salary_raw,
            "exclusions": exclusions,
            "exclusions_raw": exclusions_raw,
        },
        "sources": {
            # Google's public careers endpoints stopped returning a usable payload;
            # left in the code and off by default. Set true to try it again.
            "google_careers": False,
            "ats_boards": ats_boards,
            "max_pages_per_query": 2,
            "posted_within_days": 30,
            "request_timeout_seconds": 20,
            "polite_delay_seconds": 2.5,
        },
        "scoring": {
            "min_fit_score": min_fit,
            "max_results": max_results,
            "weights": {"keywords": 55, "title": 30, "experience": 15},
        },
        "delivery": {
            "to_email": to_email,
            "subject_prefix": "Job Sweep",
            "send_empty_digest": False,
        },
        "schedule": {
            "local_time": f"{hour:02d}:{minute:02d}",
            "timezone": tz_name,
            "frequency": "weekdays" if weekdays_only else "daily",
            "cron_utc": cron,
        },
    }
    write_config(cfg)
    write_log()
    write_workflow(cron, local_desc)

    header("DONE")
    print(
        f"""Next three commands:

  1. Test locally (no email sent):
       {c('python sweep_public.py --dry-run', BOLD)}

  2. Commit and push:
       {c('git add -A && git commit -m "configure job sweep" && git push', BOLD)}

  3. In GitHub -> Settings -> Secrets and variables -> Actions, add:
       GMAIL_USER            your.address@gmail.com
       GMAIL_APP_PASSWORD    16-character Google App Password (not your login password)

Then open the Actions tab and hit "Run workflow" on Daily Job Sweep to prove it end to end.
Generate the PDF manual any time with: {c('python generate_guide_pdf.py', BOLD)}
"""
    )
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except KeyboardInterrupt:
        print("\nCancelled. Nothing was written.")
        sys.exit(130)
